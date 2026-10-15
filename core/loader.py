from __future__ import annotations
import re
from pathlib import Path

import openpyxl

from .roadmap import Course
from .prereq import normalize_name

_CHOICE =re .compile (r"(?<!선)택\s*1")
_CODE =re .compile (r"[A-Z]{3,5}\d{3}")
_SMALL_GROUP =3

def _col (header :list [str ],*keys :str )->int |None :
    for i ,h in enumerate (header ):
        if any (k in h for k in keys ):
            return i
    return None

def _credits (val )->int :

    if val is None :
        return 3
    m =re .search (r"\d+",str (val ))
    return int (m .group ())if m else 3

def _offered (val )->tuple [int ,...]|None :
    if val is None :
        return None
    sems =sorted ({int (d )for d in re .findall (r"[12]",str (val ))})
    return tuple (sems )or None

def _prefix (code :str )->str :
    m =re .match (r"[A-Z]+",code )
    return m .group ()if m else code

def _assign_groups (recs :list [dict ],explicit :bool )->None :

    if explicit :
        prev =None
        for r in recs :
            if not r ["mark"]:
                prev =None
                continue
            if (prev is not None and prev ["ctype"]==r ["ctype"]
            and prev ["idx"]==r ["idx"]-1
            and _prefix (r ["code"])!=prev ["anchor"]):
                r ["group"],r ["anchor"]=prev ["group"],prev ["anchor"]
            else :
                r ["group"]=f"택1:{r ['ctype']}:{r ['code']}"
                r ["anchor"]=_prefix (r ["code"])
            prev =r
        return
    by_type :dict [str ,list [dict ]]={}
    for r in recs :
        if _CHOICE .search (r ["ctype"]):
            by_type .setdefault (r ["ctype"],[]).append (r )
    for ctype ,members in by_type .items ():
        if 2 <=len (members )<=_SMALL_GROUP :
            for r in members :
                r ["group"]=f"택1:{ctype }"

def load_major (xlsx :str |Path ,sheet :str )->list [Course ]:

    wb =openpyxl .load_workbook (xlsx ,data_only =True )
    ws =wb [sheet ]
    rows =list (ws .iter_rows (values_only =True ))
    if not rows :
        return []
    header =[str (c ).strip ()if c else ""for c in rows [0 ]]

    i_type =_col (header ,"구분")
    i_code =_col (header ,"학수번호")
    i_name =_col (header ,"교과목명")
    i_cred =_col (header ,"학점")
    i_rec =_col (header ,"권장이수")
    i_open =_col (header ,"개설학기")
    i_grp =_col (header ,"택1")
    i_pre =_col (header ,"선수")

    def cell (r ,i ):
        return r [i ]if i is not None and i <len (r )else None

    recs :list [dict ]=[]
    seen_codes :set [str ]=set ()
    for idx ,r in enumerate (rows [1 :]):
        if not r or not any (r ):
            continue
        code =cell (r ,i_code )
        name =cell (r ,i_name )
        if not name :
            continue
        code_s =str (code ).strip ().upper ()if code else ""

        if not re .match (r"^[A-Z]{2,5}\d{3}",code_s ):
            base =re .sub (r"[^0-9A-Z가-힣]","",str (name ).upper ())[:12 ]or f"ROW{idx }"
            code_s =f"TBD-{base }"
            n =2
            while code_s in seen_codes :
                code_s =f"TBD-{base }-{n }"
                n +=1
        seen_codes .add (code_s )
        cred =cell (r ,i_cred )
        pre =cell (r ,i_pre )
        recs .append ({
        "idx":idx ,"code":code_s ,"name":str (name ).strip (),
        "ctype":str (cell (r ,i_type )or "전공선택"),
        "credits":_credits (cred ),
        "credits_known":cred is not None and bool (re .search (r"\d",str (cred ))),
        "rec":str (cell (r ,i_rec )).strip ()if cell (r ,i_rec )else None ,
        "offered":_offered (cell (r ,i_open )),
        "prereqs":tuple (x for x in _CODE .findall (str (pre ).upper ())if x !=code_s )
        if pre else (),
        "prereq_note":str (pre ).strip ()if pre else None ,
        "mark":bool (cell (r ,i_grp )),
        "group":None ,
        })

    _assign_groups (recs ,explicit =i_grp is not None )
    return [Course (
    code =x ["code"],name =x ["name"],credits =x ["credits"],
    course_type ="전공필수"if "필수"in x ["ctype"]else "전공선택",
    recommended =x ["rec"],choice_group =x ["group"],offered =x ["offered"],
    prereqs =x ["prereqs"],prereq_note =x ["prereq_note"],
    credits_known =x ["credits_known"],
    )for x in recs ]

def build_name_index (xlsx :str |Path )->dict [str ,str ]:

    wb =openpyxl .load_workbook (xlsx ,data_only =True )
    index :dict [str ,str ]={}
    for sheet in wb .sheetnames :
        if "교과목"not in sheet :
            continue
        ws =wb [sheet ]
        rows =list (ws .iter_rows (values_only =True ))
        if not rows :
            continue
        header =[str (c ).strip ()if c else ""for c in rows [0 ]]
        i_code =_col (header ,"학수번호")
        i_name =_col (header ,"교과목명")
        if i_code is None or i_name is None :
            continue
        for r in rows [1 :]:
            if not r or i_code >=len (r )or i_name >=len (r ):
                continue
            if r [i_code ]and r [i_name ]:
                index .setdefault (
                normalize_name (r [i_name ]),str (r [i_code ]).strip ().upper ()
                )
    return index

def list_majors (xlsx :str |Path )->list [str ]:
    wb =openpyxl .load_workbook (xlsx ,data_only =True )
    return [s for s in wb .sheetnames if "교과목"in s ]
