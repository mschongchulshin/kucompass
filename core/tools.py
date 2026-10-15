from __future__ import annotations
import json
import re
from functools import lru_cache
from pathlib import Path

from .loader import load_major ,build_name_index ,list_majors
from .prereq import PrereqGraph ,all_flowmap_edges
from .roadmap import build_roadmap as _build_roadmap ,verify_no_violation
from .requirements import check_requirements as _check_req ,get_rule

DATA =Path (__file__ ).resolve ().parent .parent /"data"
XLSX =DATA /"majors.xlsx"
EDGES =DATA /"edges.json"

@lru_cache (maxsize =1 )
def _name_index ()->dict [str ,str ]:
    return build_name_index (XLSX )

@lru_cache (maxsize =1 )
def _flow_edges ()->tuple :
    return tuple (all_flowmap_edges (EDGES ))

@lru_cache (maxsize =64 )
def _sheet_for (major_code :str )->str |None :

    import re as _re
    sheets =list_majors (XLSX )
    if not major_code :
        return None

    def norm (x :str )->str :
        return _re .sub (r"[\s()·&\-_]|융합전공|전공|교과목","",str (x )).lower ()

    key =norm (major_code )
    if not key :
        return None

    for s_ in sheets :
        if norm (s_ )==key :
            return s_
    for s_ in sheets :
        ns =norm (s_ )
        if key in ns or ns in key :
            return s_

    aliases ={
    "aice":"인공지능응용융합전공_교과목",
    "aiap":"인공지능응용융합전공_교과목",
    "aiapp":"인공지능응용융합전공_교과목",
    "ai":"인공지능응용융합전공_교과목",
    "인공지능응용":"인공지능응용융합전공_교과목",
    "인공지능":"인공지능응용융합전공_교과목",
    "정보보호":"정보보호융합전공_교과목",
    "보안":"정보보호융합전공_교과목",
    "뇌인지":"뇌인지과학융합전공_교과목",
    "금융공학":"금융공학융합전공_교과목",
    "소프트웨어기술벤처":"소프트웨어기술벤처융합전공_교과목",
    "소프트웨어":"소프트웨어기술벤처융합전공_교과목",
    "swtv":"소프트웨어기술벤처융합전공_교과목",
    }
    for k ,v in aliases .items ():
        if (k in key or key in k )and v in sheets :
            return v

    if len (key )>=2 :
        for s_ in sheets :
            if norm (s_ ).startswith (key [:2 ]):
                return s_
    return None

@lru_cache (maxsize =64 )
def _graph_for (sheet :str ,enable_l3 :bool =False )->PrereqGraph :
    courses =load_major (XLSX ,sheet )
    g =PrereqGraph ({c .code :c .name for c in courses },enable_l3 =enable_l3 )
    g .add_flowmap_edges (list (_flow_edges ()),_name_index ())
    g .add_sheet_prereqs ({c .code :c .prereqs for c in courses if c .prereqs })
    g .add_sequence_rules ()
    g .add_code_rules ()
    return g

def _as_int (value ,default :int )->int :

    try :
        return int (str (value ).strip ())
    except (TypeError ,ValueError ):
        return default

def _split_codes (value )->tuple [list [str ],list [str ]]:

    import re as _re
    if value is None :
        return [],[]
    if isinstance (value ,str ):
        parts =[p for p in _re .split (r"[,\s]+",value )if p .strip ()]
    elif isinstance (value ,(list ,tuple ,set )):
        parts =[str (p )for p in value ]
    else :
        parts =[str (value )]
    out ,unknown =[],[]
    for raw in parts :
        p =raw .strip ().upper ()
        if not p :
            continue
        if _re .match (r"^[A-Z]{3,5}\d{3}$",p ):
            out .append (p )
        else :
            code =_code_by_name (p )
            if code :
                out .append (code )
            else :
                unknown .append (raw .strip ())
    return out ,unknown

def _as_codes (value )->list [str ]:

    return _split_codes (value )[0 ]

@lru_cache (maxsize =1 )
def _name_to_code ()->dict :
    from .prereq import normalize_name
    return {normalize_name (k ):v for k ,v in _name_index ().items ()}

def _code_by_name (text :str )->str |None :
    from .prereq import normalize_name
    return _name_to_code ().get (normalize_name (text ))

@lru_cache (maxsize =1 )
def _official_names ()->dict :

    import openpyxl
    import re as _re
    wb =openpyxl .load_workbook (XLSX ,data_only =True )
    official =[str (r [0 ]).strip ()for r in
    list (wb ["융합전공_목록"].iter_rows (values_only =True ))[1 :]
    if r and r [0 ]]

    def norm (x ):
        t =_re .sub (r"[\s()&·]|융합전공|전공|교과목","",str (x )).lower ()

        return t .replace ("쳐","처").replace ("됴","도")

    out ={}
    for sheet in wb .sheetnames :
        if "교과목"not in sheet :
            continue
        ns =norm (sheet )
        best =None
        for name in official :
            no =norm (name )
            if no and (no ==ns or no in ns or ns in no ):
                if best is None or len (no )>len (norm (best )):
                    best =name
        if best is None :

            best =sheet .replace ("_교과목","").replace ("융합전공","")
            best =best .replace ("_폐지"," (폐지)")or sheet
        out [sheet ]=best
    return out

def _display_name (sheet :str )->str :
    return _official_names ().get (sheet ,sheet .replace ("_교과목",""))

_DEPT_SUFFIX =("학부","학과","대학원","대학","전공","부")

def _normalize_dept (name :str )->str :
    n =re .sub (r"\s+|·|‧|\(.*?\)","",str (name ))
    for suf in _DEPT_SUFFIX :
        if n .endswith (suf )and len (n )>len (suf )+1 :
            n =n [:-len (suf )]
            break
    return n

@lru_cache (maxsize =1 )
def _known_departments ()->tuple [str ,...]:

    import openpyxl
    out =set ()
    try :
        wb =openpyxl .load_workbook (XLSX ,read_only =True )
        for sheet ,cols in (("융합전공_목록",(1 ,2 )),("이중전공_선발제외학과",(0 ,))):
            if sheet not in wb .sheetnames :
                continue
            ws =wb [sheet ]
            for r in list (ws .iter_rows (values_only =True ))[1 :]:
                for col in cols :
                    if col <len (r )and r [col ]:
                        for part in str (r [col ]).split ("/"):
                            n =_normalize_dept (part )
                            if len (n )>=2 :
                                out .add (n )
    except Exception :
        return ()
    return tuple (sorted (out ))

def _looks_like_department (name :str )->bool :

    n =_normalize_dept (name )
    if len (n )<2 :
        return False
    for d in _known_departments ():
        if n ==d or (len (n )>=3 and (n in d or d in n )):
            return True
    return False

def _eun_neun (word :str )->str :

    last =str (word ).strip ()[-1 :]or "x"
    if not ("가"<=last <="힣"):
        return "는"
    return "은"if (ord (last )-0xAC00 )%28 else "는"

def _active_major_count ()->int :

    return sum (1 for v in _official_names ().values ()if "폐지"not in str (v ))

def _not_found (major_code :str )->dict :
    if _looks_like_department (major_code ):
        msg =(f"'{major_code }'{_eun_neun (major_code )} 학과·주전공으로 보입니다. "
        f"이 시스템은 융합전공 {_active_major_count ()}개만 다루며 "
        "주전공·이중전공 커리큘럼은 포함하고 있지 않습니다. "
        "관심 분야를 말씀해 주시면 맞는 융합전공을 찾아드리겠습니다.")
    else :
        msg =(f"'{major_code }' 융합전공을 찾을 수 없습니다. "
        "전공명을 다시 확인해 주시거나, 관심 분야를 말씀해 주시면 "
        "후보를 찾아드리겠습니다.")
    return {"error":msg ,"display_text":msg }

_SEMESTER =re .compile (r"^\s*([1-4])\s*(?:/|-|학년\s*)\s*([12])\s*(?:학기)?\s*$")
_SCOPE =re .compile (r"이중\s*전공|복수\s*전공|주\s*전공|부\s*전공|전과")
SCOPE_NOTE =("이 서비스는 고려대 서울캠퍼스 융합전공만 다룹니다. "
"주전공·이중전공·부전공 교육과정은 포함하고 있지 않으니 소속 학과나 학사 공지에서 확인해 주세요.")

def asks_out_of_scope (text :str )->bool :
    t =str (text or "")
    return bool (_SCOPE .search (t ))and "융합"not in t

def _norm_semester (value )->str |None :
    m =_SEMESTER .match (str (value or ""))
    return f"{m .group (1 )}/{m .group (2 )}"if m else None

def _norm_course_type (value )->str |None :
    t =re .sub (r"\s","",str (value or ""))
    if "필수"in t :
        return "전공필수"
    if "선택"in t :
        return "전공선택"
    return None

@lru_cache (maxsize =1 )
def _known_codes ()->frozenset :
    return frozenset (_name_index ().values ())

def _abolished_notice (name :str )->str |None :
    if "폐지"not in name :
        return None
    return (f"{name }{_eun_neun (name )} 폐지된 융합전공입니다. 새로 이수할 수 있는지는 학과에 확인하세요. "
    "아래는 기존 이수자 기준 계산입니다.")

def _no_credit_info (sheet :str ,courses :list )->dict |None :
    if any (c .credits_known for c in courses ):
        return None
    import openpyxl
    ws =openpyxl .load_workbook (XLSX ,read_only =True ,data_only =True )[sheet ]
    areas :list [str ]=[]
    for r in list (ws .iter_rows (values_only =True ))[1 :]:
        if r and r [0 ]and str (r [0 ]).strip ()not in areas :
            areas .append (str (r [0 ]).strip ())
    name =_display_name (sheet )
    msg =(f"{name }{_eun_neun (name )} 학교 자료에 과목별 학점과 학수번호가 공개되지 않아 "
    "로드맵과 졸업요건을 계산할 수 없습니다. "
    f"영역별 요건은 {', '.join (areas )}입니다. "
    "교과목 목록은 교과목 조회로 볼 수 있고, 이수 계획은 주관 학과에 확인해 주세요.")
    return {"error":msg ,"display_text":msg ,"major_name":name ,"data_insufficient":True }

def _course_line (c )->str :
    code ="학수번호 미공개"if c .code .startswith ("TBD")else c .code
    cred =f"{c .credits }학점"if c .credits_known else "학점 미공개"
    when =(f", 권장 {c .recommended }"if c .recommended else "")+(
    f", 개설 {'·'.join (map (str ,c .offered ))}학기"if c .offered else "")
    return f"{c .course_type } {code } {c .name } ({cred }{when })"

def query_curriculum (major_code :str ,course_type :str |None =None ,
year_semester :str |None =None )->dict :
    """전공의 교과목 목록 조회. 과목, 커리큘럼, 전공필수 과목을 물으면 호출.
    기억에 의존하지 말고 이 결과만 사용하라."""
    sheet =_sheet_for (major_code )
    if not sheet :
        return _not_found (major_code )
    courses =load_major (XLSX ,sheet )
    total_unfiltered =len (courses )
    ctype =_norm_course_type (course_type )
    if ctype :
        courses =[c for c in courses if c .course_type ==ctype ]

    semester_note =None
    ys =_norm_semester (year_semester )if year_semester else None
    if ys :
        if any (c .recommended for c in courses ):
            picked =[c for c in courses if c .recommended ==ys ]
            if picked :
                courses =picked
            else :
                semester_note =f"권장학기가 {ys }인 과목이 없어 전체를 보여드립니다."
                ys =None
        else :
            semester_note =(f"이 전공은 권장학기 정보가 없어 "
            f"'{ys }' 학기 필터를 적용하지 않았습니다.")
            ys =None

    items =[{"code":c .code ,"name":c .name ,"type":c .course_type ,
    "credits":c .credits if c .credits_known else None }for c in courses [:15 ]]
    lines =[_course_line (c )for c in courses [:20 ]]
    more =f"\n... 외 {len (courses )-20 }과목"if len (courses )>20 else ""

    if not courses and total_unfiltered :
        empty =(f"'{ctype }' 구분의 과목이 없습니다. "
        f"(전체 {total_unfiltered }과목)")
    else :
        empty =None

    out ={"major_name":_display_name (sheet ),
    "courses_sample":items ,"total_count":len (courses ),
    "total_unfiltered":total_unfiltered ,
    "filter_applied":ctype or ys ,
    "display_text":"\n".join (lines )+more }
    for k ,v in (("semester_note",semester_note ),("empty_reason",empty )):
        if v :
            out [k ]=v
            out ["display_text"]=(v +"\n"+out ["display_text"]).strip ()
    return out

def check_requirements (major_code :str ,taken_courses :list [str ]|None =None ,
first_major :str |None =None )->dict :
    """졸업요건 충족 여부와 잔여 학점 계산. 졸업요건, 학점, 더 들어야 할 과목을 물으면 호출.
    taken_courses에는 학생이 들었다고 말한 과목을 그대로 넣는다.
    과목명만 말했으면 과목명을 넣고 학수번호를 지어내지 마라."""
    sheet =_sheet_for (major_code )
    if not sheet and first_major :

        sheet =_sheet_for (first_major )
    if not sheet :
        return _not_found (major_code )
    courses =load_major (XLSX ,sheet )
    bad =_no_credit_info (sheet ,courses )
    if bad :
        return bad
    rule =get_rule (sheet ,courses )
    name =_display_name (sheet )
    codes ,unknown_names =_split_codes (taken_courses )
    res =_check_req (courses ,codes ,rule ,name ,known_codes =set (_known_codes ()),
    unknown_names =unknown_names )
    notes =res .notes ()
    notice =_abolished_notice (name )
    text =res .display_text ()
    return {
    "major_name":name ,
    "total_required":res .total_required ,"core_required":res .core_required ,
    "earned_total":res .earned_total ,"earned_core":res .earned_core ,
    "remaining_total":res .remaining_total ,"remaining_core":res .remaining_core ,
    "overlap_allowed":res .overlap_allowed ,"satisfied":res .satisfied ,
    "unmet":res .unmet ,"rule_note":rule .note ,
    "notes":([notice ]if notice else [])+notes ,
    "missing_core":res .missing_core ,
    "display_text":f"{notice } {text }"if notice else text ,
    }

def build_roadmap (major_code :str ,taken_courses :list [str ]|None =None ,
start_year_semester :str ="3/1",
max_credits_per_term :int =18 )->dict :
    """선수과목 순서를 반영한 학기별 수강 로드맵 생성. 수강 계획, 로드맵, 학기별 순서를 물으면 호출.
    학기당 학점이나 현재 학기는 학생이 말한 경우에만 넣고, 이수 과목은 학생이 말한 그대로 넣는다."""
    sheet =_sheet_for (major_code )
    if not sheet :
        return _not_found (major_code )
    courses =load_major (XLSX ,sheet )
    bad =_no_credit_info (sheet ,courses )
    if bad :
        return bad
    graph =_graph_for (sheet )
    rule =get_rule (sheet ,courses )
    taken =set (_as_codes (taken_courses ))

    ys =_norm_semester (start_year_semester )or "3/1"
    y ,s =ys .split ("/")
    start =(int (y )-1 )*2 +int (s )
    cap =max (3 ,min (24 ,_as_int (max_credits_per_term ,18 )))

    rm =_build_roadmap (courses ,graph ,taken =taken ,start_order =start ,
    max_credits =cap ,
    target_credits =rule .total_required ,
    required_core =rule .core_required )
    violations =verify_no_violation (rm ,graph ,taken )
    name =_display_name (sheet )
    notice =_abolished_notice (name )

    return {
    "major_name":name ,
    "terms":[{"term":t .label ,"term_credits":t .credits ,
    "courses":[c .code for c in t .courses ],
    "names":[c .label for c in t .courses ]}for t in rm .terms ],
    "total_terms":len (rm .terms ),"total_credits":rm .total_credits ,
    "taken_credits":rm .taken_credits ,
    "target_credits":rule .total_required ,
    "warnings":([notice ]if notice else [])+rm .warnings ,
    "prereq_violations":len (violations ),
    "completed_early":rm .completed_early ,
    "unplaced_count":len (rm .unplaced ),
    "edges_used":rm .edges_used ,
    "display_text":rm .display_text (),
    }

def search_majors (interest_text :str ,campus :str |None =None ,
top_k :int =5 )->dict :
    """관심사 자연어로 융합전공 후보 검색. 학생이 관심 분야나 진로를 말하며 맞는 융합전공을 물을 때만 호출.
    학생이 이미 전공 이름을 말했으면 호출하지 말고 그 전공으로 다른 툴을 바로 호출하라."""
    from .recommend import recommend
    r =recommend (interest_text ,campus =campus ,top_k =top_k )
    if asks_out_of_scope (interest_text ):
        r =dict (r )
        r ["scope_note"]=SCOPE_NOTE
        r ["display_text"]=SCOPE_NOTE +"\n\n"+r ["display_text"]
    return r

def get_syllabus (course_code :str ,year :int =2024 ,term :int =2 )->dict :
    """개별 과목의 강의계획서(주차별 내용, 평가방식) 조회. 강의계획서를 물을 때만 호출.
    오프라인 모드에서는 조회되지 않을 수 있으며, 결과에 없는 내용은 지어내지 마라."""
    from .syllabus import fetch_syllabus
    return fetch_syllabus (course_code ,_as_int (year ,2024 ),_as_int (term ,2 ))

TOOLS =[search_majors ,query_curriculum ,check_requirements ,
build_roadmap ,get_syllabus ]
